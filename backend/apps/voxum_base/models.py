from openpyxl import load_workbook
from pathlib import PurePath
from django.contrib.contenttypes.fields import GenericForeignKey
from django.contrib.contenttypes.models import ContentType
from django.db import models
from utils import _

from core.abstract.models import AbstractModel


class AbstractInfo(AbstractModel):
    legal_number = models.CharField('CPF/CNPJ', max_length=18, unique=True, db_index=True)

    class Meta:
        abstract = True
        ordering = ('-created_at', '-updated_at')

    def __str__(self):
        return f'{self.legal_number}'


class Entity(AbstractInfo):
    """Class responsible for details to credores"""


CHOICES_COIN = (
    ("B", "BRL"),
)


class Coin(AbstractModel):
    type = models.CharField(max_length=1, verbose_name=_('Coin'), choices=CHOICES_COIN, default='B')

    def __str__(self):
        return f"{self.type} - {self.get_type_display()}"


class ExcelWorker(AbstractModel):
    STATUS_PENDING = 'P'
    STATUS_SUCCEEDED = 'S'
    STATUS_FAILED = 'F'
    STATUS_CHOICES = (
        (STATUS_PENDING, _('Pendente')),
        (STATUS_SUCCEEDED, _('Concluído')),
        (STATUS_FAILED, _('Falhou')),
    )

    file = models.FileField(upload_to='process_files/%Y/%m/%d/')
    content_type = models.ForeignKey(ContentType, on_delete=models.CASCADE)
    object_id = models.UUIDField()
    content_object = GenericForeignKey('content_type', 'object_id')
    status = models.CharField(max_length=1, choices=STATUS_CHOICES, default=STATUS_PENDING)

    def process(self, callback):
        if not self.file:
            raise ValueError('No spreadsheet was supplied for processing.')

        with self.file.open('rb') as source:
            if PurePath(self.file.name).suffix.lower() == '.xls':
                import xlrd

                worksheet = xlrd.open_workbook(file_contents=source.read()).sheet_by_index(0)
                headers = worksheet.row_values(0)
                rows = (worksheet.row_values(index) for index in range(1, worksheet.nrows))
            else:
                workbook = load_workbook(source, read_only=True, data_only=True)
                worksheet = workbook.active
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, ())

            columns = [str(value).strip() if value is not None else '' for value in headers]
            data = [
                {**dict(zip(columns, row)), 'INDEX': index}
                for index, row in enumerate(rows, start=2)
            ]

        try:
            callback(
                data,
                instance_id=str(self.pk),
                object_id=str(self.object_id),
            )
        except Exception:
            self.status = self.STATUS_FAILED
            self.save(update_fields=('status', 'updated_at'))
            raise

        self.status = self.STATUS_SUCCEEDED
        self.save(update_fields=('status', 'updated_at'))


class ErrorFile(AbstractModel):
    STATUS_PENDING = 'P'
    STATUS_REVIEWED = 'R'
    STATUS_RESOLVED = 'S'
    STATUS_CHOICES = (
        (STATUS_PENDING, _('Pendente')),
        (STATUS_REVIEWED, _('Revisado')),
        (STATUS_RESOLVED, _('Resolvido')),
    )

    file = models.ForeignKey(ExcelWorker, on_delete=models.CASCADE, related_name='errors')
    error = models.TextField()
    data = models.JSONField(default=dict, blank=True)
    status = models.CharField(max_length=1, choices=STATUS_CHOICES, default=STATUS_PENDING)
    traceback = models.TextField(blank=True)